import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from digest import store
from digest.extract_api import extract_pdf, merge_continuations
from digest.pdf_pages import HEADER_RE, PageInfo, guess_date, parse_page_range
from digest.report import build_cumulative_html, build_cumulative_xlsx, build_daily_docx
from digest.schema import ROOT, load_categories, make_id, normalize_record, response_schema, category_names

CATS = load_categories()
NAMES = category_names(CATS)
SAMPLE = ROOT / "samples" / "2026-10-08.sample.json"
REAL_PDF = os.environ.get("PRESS_DIGEST_TEST_PDF")  # 실제 PDF가 있을 때만 통합 테스트 실행


def rec(**kw):
    base = dict(media="매일경제", title="테스트 기사", page_label="A01면", category="은행", sub_categories=[],
                issue="", relevance="높음", facts=["수치 1조원"], situation="s", stances=[], implications="i", notes="")
    base.update(kw)
    return base


class HeaderAndDate(unittest.TestCase):
    def test_header_regex(self):
        m = HEADER_RE.search("2026년 10월 8일 목요일 B02면 경제 금융")
        self.assertEqual((m["y"], m["m"], m["d"], m["rest"]), ("2026", "10", "8", "B02면 경제 금융"))

    def test_guess_date_priority(self):
        pages = [PageInfo(2, ["A06면"], ["2026-10-08"]), PageInfo(3, ["A06면"], ["2026-10-08"])]
        self.assertEqual(guess_date("x_261009.pdf", pages), "2026-10-08")       # 헤더 > 파일명
        self.assertEqual(guess_date("x_261009.pdf", pages, "2026-01-02"), "2026-01-02")  # 지정 > 헤더
        self.assertEqual(guess_date("보도_261009.pdf", []), "2026-10-09")        # 파일명 폴백
        with self.assertRaises(ValueError):
            guess_date("none.pdf", [])

    def test_page_range(self):
        self.assertEqual(parse_page_range("2-4,9", 59), [2, 3, 4, 9])
        self.assertEqual(parse_page_range(None, 3), [1, 2, 3])


class Schema(unittest.TestCase):
    def test_normalize_unknown_category_falls_back(self):
        r = normalize_record(rec(category="없는분야", sub_categories=["은행", "없음"]), "2026-10-08", NAMES)
        self.assertEqual(r["category"], "기타·타부처")
        self.assertEqual(r["sub_categories"], ["은행"])

    def test_id_is_stable_across_spacing_and_punctuation(self):
        self.assertEqual(make_id("2026-10-08", "한겨레", "제목, 입니다!"), make_id("2026-10-08", "한겨레", "제목 입니다"))

    def test_schema_is_strict(self):
        art = response_schema(NAMES)["properties"]["articles"]["items"]
        self.assertFalse(art["additionalProperties"])
        self.assertEqual(set(art["required"]), set(art["properties"]))


class Store(unittest.TestCase):
    def test_upsert_dedup_and_query(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            a = normalize_record(rec(), "2026-10-08", NAMES, pdf_page=2)
            b = normalize_record(rec(title="다른 기사", category="보험", issue="X"), "2026-10-08", NAMES, pdf_page=3)
            self.assertEqual(store.upsert([a, b], "2026-10-08", base), (2, 0))
            a2 = dict(a, situation="수정됨")
            self.assertEqual(store.upsert([a2], "2026-10-08", base), (0, 1))   # 재처리는 덮어쓰기
            allr = store.load_all(base)
            self.assertEqual(len(allr), 2)
            self.assertEqual([r["title"] for r in store.query(allr, category="보험")], ["다른 기사"])
            self.assertEqual(len(store.query(allr, q="1조원")), 2)
            self.assertEqual(store.query(allr, since="2026-10-09"), [])


class FakeClient:
    """messages.create 를 흉내내 JSON 문자열을 돌려주는 가짜 클라이언트."""
    def __init__(self, per_page):
        self.per_page, self.calls = per_page, []
        self.messages = self

    def create(self, **kw):
        self.calls.append(kw)
        txt = kw["messages"][0]["content"][0]["text"]
        page = int(txt.split("PDF ")[1].split("페이지")[0])
        if page == 99:
            return SimpleNamespace(stop_reason="refusal", content=[])
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text",
                               text=json.dumps({"articles": self.per_page[page]}, ensure_ascii=False))])


@unittest.skipUnless(REAL_PDF, "PRESS_DIGEST_TEST_PDF 미설정")
class RealPdf(unittest.TestCase):
    def test_pages_and_fake_extract(self):
        from digest.pdf_pages import read_pages, render_tiles
        pages = read_pages(REAL_PDF)
        self.assertEqual(guess_date(REAL_PDF, pages), "2026-10-08")
        self.assertFalse(pages[0].headers)                       # 표지는 헤더 없음 → 처리 제외 대상
        self.assertTrue(all(p.headers for p in pages[1:]))
        tiles = render_tiles(REAL_PDF, 3)
        self.assertEqual(len(tiles), 2)
        self.assertTrue(all(t[:2] == b"\xff\xd8" and len(t) < 5_000_000 for t in tiles))
        art = dict(rec(), continues_previous=False)
        client = FakeClient({3: [art]})
        recs, failed = extract_pdf(REAL_PDF, [pages[2]], "2026-10-08", CATS, client=client, log=lambda *_: None)
        self.assertEqual((len(recs), failed), (1, []))
        content = client.calls[0]["messages"][0]["content"]
        self.assertEqual(sum(c["type"] == "image" for c in content), 2)
        self.assertEqual(client.calls[0]["output_config"]["format"]["type"], "json_schema")
        self.assertNotIn("tool_choice", client.calls[0])          # Opus 5.5 는 강제 tool_choice 불가


class Merge(unittest.TestCase):
    def test_continuation_merges_into_previous_same_media(self):
        a = dict(rec(), continues_previous=False, pdf_page=4, stances=[])
        b = dict(rec(title="", facts=["추가 수치 5%"], situation="", implications=""), continues_previous=True, pdf_page=5, stances=[])
        c = dict(rec(media="한겨레", title="다른 매체"), continues_previous=False, pdf_page=6, stances=[])
        out = merge_continuations([(5, [b]), (4, [a]), (6, [c])])   # 순서가 뒤섞여 들어와도 페이지순 처리
        self.assertEqual(len(out), 2)
        self.assertEqual(out[0]["facts"], ["수치 1조원", "추가 수치 5%"])
        self.assertIn("p.5 이어짐 병합", out[0]["notes"])


class Reports(unittest.TestCase):
    def test_reports_from_sample(self):
        raw = json.loads(SAMPLE.read_text(encoding="utf-8"))
        recs = [normalize_record(r, "2026-10-08", NAMES) for r in raw]
        self.assertEqual(len({r["id"] for r in recs}), len(recs))
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            from docx import Document
            from openpyxl import load_workbook
            docx = build_daily_docx(recs, "2026-10-08", CATS, d / "a.docx")
            doc = Document(docx)
            self.assertEqual(doc.tables[0].rows[0].cells[3].text, "사실관계(수치)")
            self.assertEqual([c.text for c in doc.tables[0].rows[0].cells][:3], ["일자", "매체명", "제목"])
            self.assertEqual(sum(len(t.rows) - 1 for t in doc.tables), len(recs))
            wb = load_workbook(build_cumulative_xlsx(recs, CATS, d / "a.xlsx"))
            self.assertEqual(wb["전체"].max_row, len(recs) + 1)
            self.assertIn("은행", wb.sheetnames)
            html = build_cumulative_html(recs, CATS, d / "a.html").read_text(encoding="utf-8")
            self.assertNotIn("__DATA__", html)
            self.assertNotIn("</script><", html.split('id="data"')[1].split("</script>")[0])  # JSON 내 스크립트 종료 방지


if __name__ == "__main__":
    unittest.main()
