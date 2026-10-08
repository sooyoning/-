"""보고서 생성: 일일 DOCX, 누적 XLSX, 누적 HTML(필터·그룹 보기)."""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import STANCE_GROUPS, category_names

# 요청된 정리 순서: 일자 - 매체명 - 제목 - 사실관계(수치) - 상황 - 금융위·금감원 등 및 이해관계자 입장 - 시사점
COLUMNS = ["일자", "매체명", "제목", "사실관계(수치)", "상황", "금융위·금감원 등 및 이해관계자 입장", "시사점"]
REL_RANK = {"높음": 0, "보통": 1, "낮음": 2}
WEEKDAY = "월화수목금토일"


def fmt_date(d: str) -> str:
    y, m, dd = map(int, d.split("-"))
    return f"{y}.{m:02d}.{dd:02d}({WEEKDAY[dt.date(y, m, dd).weekday()]})"


def facts_text(r: dict) -> str:
    return "\n".join(f"• {f}" for f in r["facts"])


def stances_text(r: dict) -> str:
    order = {g: i for i, g in enumerate(STANCE_GROUPS)}
    sts = sorted(r["stances"], key=lambda s: order.get(s["group"], 99))
    return "\n".join(f"• [{s['group']}] {s['actor']}: {s['position']}" if s["actor"]
                     else f"• [{s['group']}] {s['position']}" for s in sts)


def row_values(r: dict) -> list[str]:
    return [fmt_date(r["date"]), r["media"], r["title"], facts_text(r), r["situation"],
            stances_text(r), r["implications"]]


def _sort_key(cat_order: dict[str, int]):
    return lambda r: (cat_order.get(r["category"], 99), r["issue"], REL_RANK[r["relevance"]], r["media"])


# ---------------------------------------------------------------- 일일 DOCX
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_font(run, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _fill_cell(cell, text: str, size: float = 8, bold: bool = False) -> None:
    cell.text = ""
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _set_font(p.add_run(line), size, bold)


def build_daily_docx(records: list[dict], date: str, cats: list[dict], path: Path,
                     min_relevance: str | None = None) -> Path:
    """분야별로 묶은 일일 표. min_relevance='보통' 이면 '낮음' 기사는 제외."""
    names = category_names(cats)
    order = {n: i for i, n in enumerate(names)}
    recs = [r for r in records if r["date"] == date]
    if min_relevance:
        recs = [r for r in recs if REL_RANK[r["relevance"]] <= REL_RANK[min_relevance]]
    recs.sort(key=_sort_key(order))

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.0))

    p = doc.add_paragraph()
    _set_font(p.add_run(f"신문기사 주요 내용 정리 — {fmt_date(date)}"), 15, True)
    p = doc.add_paragraph()
    _set_font(p.add_run(f"기사 {len(recs)}건 · 소관분야 "
                        f"{len({r['category'] for r in recs})}개 · 시사점은 기사 근거 범위 내 분석(의견)이며 사실관계와 구분됨"),
              8.5, color="666666")

    widths = [Cm(2.1), Cm(1.8), Cm(3.4), Cm(7.6), Cm(3.4), Cm(5.4), Cm(3.9)]
    for cat in names:
        group = [r for r in recs if r["category"] == cat]
        if not group:
            continue
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        _set_font(h.add_run(f"■ {cat} ({len(group)}건)"), 11, True, "1F3864")
        table = doc.add_table(rows=1, cols=len(COLUMNS))
        table.style = "Table Grid"
        table.autofit = False
        for i, w in enumerate(widths):  # LibreOffice/한글은 tblGrid 열 너비를 기준으로 렌더링한다
            table.columns[i].width = w
        for i, name in enumerate(COLUMNS):
            c = table.rows[0].cells[i]
            c.width = widths[i]
            _fill_cell(c, name, 8.5, True)
            _shade(c, "D9E2F3")
        for r in group:
            vals = row_values(r)
            if r["issue"]:
                vals[2] = f"{r['title']}\n〔이슈: {r['issue']}〕"
            cells = table.add_row().cells
            for i, v in enumerate(vals):
                cells[i].width = widths[i]
                _fill_cell(cells[i], v, 8, bold=(i == 2))
            if r["notes"]:
                _fill_cell(cells[2], vals[2] + f"\n※ {r['notes']}", 8, bold=False)
        # 헤더 행 반복
        trPr = table.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------- 누적 XLSX
def build_cumulative_xlsx(records: list[dict], cats: list[dict], path: Path) -> Path:
    names = category_names(cats)
    order = {n: i for i, n in enumerate(names)}
    head = COLUMNS + ["소관분야", "관련분야", "이슈", "관련도", "면", "PDF쪽"]

    def fill(ws, rows: list[dict]) -> None:
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E2F3")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in sorted(rows, key=lambda r: (r["date"], -order.get(r["category"], 99)), reverse=True):
            ws.append(row_values(r) + [r["category"], ", ".join(r["sub_categories"]), r["issue"],
                                       r["relevance"], r["page_label"], r["pdf_page"]])
        for i, w in enumerate([13, 10, 32, 70, 34, 50, 38, 16, 18, 22, 8, 14, 7], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws = wb.active
    ws.title = "전체"
    fill(ws, records)
    for cat in names:
        rows = [r for r in records if r["category"] == cat]
        if rows:
            fill(wb.create_sheet(cat[:28].replace("/", "·")), rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- 누적 HTML
def build_cumulative_html(records: list[dict], cats: list[dict], path: Path) -> Path:
    payload = {"categories": category_names(cats), "stanceGroups": STANCE_GROUPS,
               "records": sorted(records, key=lambda r: r["date"], reverse=True),
               "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}
    blob = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    html = (Path(__file__).parent / "template.html").read_text(encoding="utf-8").replace("__DATA__", blob)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path
